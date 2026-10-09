"""Render one frozen ReportDocument without a model call or external network."""
from __future__ import annotations

import base64
import html
import importlib.metadata
import json
from pathlib import Path
import threading
import zipfile

from .report_document import canonical, digest, safe_text
from .report_figures import render_figures

_RENDER_LOCK = threading.RLock()


def value_text(value):
    return format(value, ".12g") if isinstance(value, (float, int)) else str(value or "—")


def sections(document):
    zh = document["language"] == "zh-CN"
    identity = document["identity"]
    labels = {"training": "训练任务", "inference": "推理任务", "diagnostic": "连接与执行诊断", "analysis": "数据分析",
        "completed": "执行已完成", "running": "执行中", "paused": "已暂停", "pausing": "正在安全暂停",
        "failed": "失败", "blocked": "需要处理", "recovering": "等待恢复核对", "queued": "排队中",
        "hash_verified": "文件哈希已核对", "independently_verified": "独立验证已通过", "unverified": "尚未验证",
        "ready": "已生成", "partial": "部分生成，仍有缺失项", "higher": "越高越好", "lower": "越低越好", "none": "不适用"}
    label = lambda value: labels.get(value, value) if zh else value
    result = [
        {"title": "执行结论与边界" if zh else "Outcome and evidence boundary", "paragraphs": [
            f"Run: {identity['run_id']}", f"{'项目' if zh else 'Project'}: {identity['project_id'] or ('未关联' if zh else 'not associated')}",
            f"{'任务类型' if zh else 'Type'}: {label(document['kind'])} | {'执行' if zh else 'Execution'}: {label(document['execution_status'])} | {'证据' if zh else 'Evidence'}: {label(document['evidence_status'])}",
            f"{'报告' if zh else 'Report'}: {label(document['report_status'])}", *document["limitations"]]},
    ]
    if document.get("author_summary"):
        result.append({"title": "说明（须结合证据审阅）" if zh else "Commentary (review alongside evidence)", "paragraphs": [document["author_summary"]]})
    result.append({"title": "指标与适用口径" if zh else "Metrics and scope",
        "headers": ["指标", "数值", "单位", "指标方向", "来源"] if zh else ["Metric", "Value", "Unit", "Direction", "Source"],
        "rows": [[row["name"], value_text(row["value"]), row["unit"], label(row["direction"]), row["source_artifact_id"]]
                 for row in document["metrics"]],
        "paragraphs": [] if document["metrics"] else ["没有可用的来源绑定指标；未填入零值。" if zh else "No source-bound metrics; missing values were not replaced with zero."]})
    result.append({"title": "执行记录" if zh else "Execution records", "headers": ["Tool", "Status", "Result", "Call ID"],
        "rows": [[row["tool"], row["status"], "passed" if row["ok"] is True else "failed" if row["ok"] is False else "unknown", row["call_id"]] for row in document["checks"]]})
    result.append({"title": "来源与文件身份" if zh else "Sources and file identity", "headers": ["Artifact", "Bytes", "SHA-256", "Verification"],
        "rows": [[safe_text(row["name"], 200), str(row["bytes"]), row["sha256"], row["verification"]] for row in document["sources"]]})
    result.append({"title": "缺失项与下一步" if zh else "Missing evidence and next step", "paragraphs": document["missing"] or [
        "按权限审阅产物；报告生成成功不代替科研结论验收。" if zh else "Review the artifacts within their access scope; rendering success is not scientific acceptance."]})
    return result


def render_markdown(document, content, figures):
    parts = ["# " + document["title"], "", f"ReportDocument: `{document['document_sha256']}`", ""]
    for section in content:
        parts.extend(["## " + section["title"], "", *section.get("paragraphs", []), ""])
        if section.get("rows"):
            clean = lambda row: [str(cell).replace("|", "\\|").replace("\n", " ") for cell in row]
            table_lines = ["| " + " | ".join(clean(section["headers"])) + " |", "| " + " | ".join("---" for _ in section["headers"]) + " |"]
            table_lines.extend("| " + " | ".join(clean(row)) + " |" for row in section["rows"])
            parts.append("\n".join(table_lines))
    parts.extend(["", "## 图表 / Figures", ""])
    for figure in figures:
        parts.extend([f"### {figure['id']} · {figure['title']}", figure["caption"]])
        parts.append(f"![{figure['title']}](figures/{figure['png']})" if figure["status"] == "ready" else "Missing data: " + figure.get("reason", "unknown"))
    return "\n\n".join(parts) + "\n"


def render_html(document, content, figures, root):
    esc = lambda value: html.escape(str(value), quote=True)
    body = [f"<header><p>DeepEvo · {esc(document['kind'])}</p><h1>{esc(document['title'])}</h1><p class='identity'>{esc(document['identity']['run_id'])}</p></header>"]
    for section in content:
        body.append(f"<section><h2>{esc(section['title'])}</h2>")
        body.extend(f"<p>{esc(value).replace(chr(10), '<br>')}</p>" for value in section.get("paragraphs", []))
        if section.get("rows"):
            body.append("<table><thead><tr>" + "".join(f"<th>{esc(value)}</th>" for value in section["headers"]) + "</tr></thead><tbody>")
            body.extend("<tr>" + "".join(f"<td>{esc(cell)}</td>" for cell in row) + "</tr>" for row in section["rows"])
            body.append("</tbody></table>")
        body.append("</section>")
    for figure in figures:
        body.append(f"<figure><h2>{esc(figure['id'])} · {esc(figure['title'])}</h2>")
        if figure["status"] == "ready":
            encoded = base64.b64encode((root / "figures" / figure["png"]).read_bytes()).decode()
            body.append(f'<img alt="{esc(figure["title"])}" src="data:image/png;base64,{encoded}">')
        else:
            body.append(f"<p>Missing data: {esc(figure.get('reason', 'unknown'))}</p>")
        body.append(f"<figcaption>{esc(figure['caption'])}<br>Source: {esc(figure.get('source_artifact_id', ''))}</figcaption></figure>")
    css = """@page { size: A4; margin: 18mm; } body { color:#102823; background:#fff; font-family:sans-serif; font-size:10pt; line-height:1.5; margin:0 auto; max-width:920px; padding:28px; }
    header { border-bottom:3px solid #147b70; padding-bottom:18px; } header p { color:#147b70; } h1 { font-size:24pt; line-height:1.25; } h2 { font-size:14pt; color:#09524b; margin-top:26px; } p,td { overflow-wrap:anywhere; } table { width:100%; border-collapse:collapse; table-layout:fixed; font-size:8pt; } th,td { text-align:left; border:1px solid #cddcd8; padding:6px; vertical-align:top; } th { background:#eaf1ef; } figure { margin:24px 0; break-inside:avoid; } img { max-width:100%; height:auto; } figcaption,.identity { font-size:8pt; color:#334b47; } @media print { body { padding:0; } thead { display:table-header-group; } }"""
    return f"<!doctype html><html lang='{document['language']}'><head><meta charset='utf-8'><title>{esc(document['title'])}</title><style>{css}</style></head><body>{''.join(body)}</body></html>"


def render_docx(document, content, figures, root):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    doc = Document()
    page = doc.sections[0]
    page.page_width, page.page_height = Cm(21), Cm(29.7)
    page.top_margin = page.bottom_margin = page.left_margin = page.right_margin = Cm(1.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor.from_string("102823")
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(document["title"], 0)
    doc.add_paragraph("ReportDocument: " + document["document_sha256"])
    for section in content:
        doc.add_heading(section["title"], 1)
        for text in section.get("paragraphs", []): doc.add_paragraph(text)
        if section.get("rows"):
            table = doc.add_table(rows=1, cols=len(section["headers"]))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, section["headers"]): cell.text = str(text)
            for row in section["rows"]:
                for cell, text in zip(table.add_row().cells, row): cell.text = str(text)
    for figure in figures:
        doc.add_heading(figure["id"] + " · " + figure["title"], 1)
        if figure["status"] == "ready": doc.add_picture(str(root / "figures" / figure["png"]), width=Cm(16.5))
        else: doc.add_paragraph("Missing data: " + figure.get("reason", "unknown"))
        doc.add_paragraph(figure["caption"])
    footer = page.footer.paragraphs[0]
    footer.text = "DeepEvo | " + document["identity"]["run_id"] + " | "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.save(root / "report.docx")


def render_pdf(html_text, output):
    import fitz
    import re
    # MuPDF leaks table-cell backgrounds onto later pages after a table ends.
    # A normal block inside the cell paints only on its actual page. Keep these
    # wrappers and font changes in the PDF copy, leaving web/Word unchanged.
    pdf_css = """th { background:transparent !important; }
    .report-pdf-heading-cell { background:#eaf1ef; }
    .report-pdf-hash { font-family: Courier, monospace !important; font-size:7pt; font-variant-ligatures:none; }
    .identity { font-family: Courier, monospace !important; }"""
    pdf_html = re.sub(r"(<td>)([A-Fa-f0-9]{64})(</td>)",
                      r'\1<span class="report-pdf-hash">\2</span>\3', html_text)
    pdf_html = re.sub(r'<th>(.*?)</th>', r'<th><div class="report-pdf-heading-cell">\1</div></th>', pdf_html, flags=re.S)
    # Keep the override in the embedded stylesheet used by the HTML renderer.
    pdf_html = pdf_html.replace('</style>', pdf_css+'</style>', 1)
    paper = fitz.paper_rect("a4")
    area = paper + (38, 38, -38, -38)
    story = fitz.Story(html=pdf_html, user_css=pdf_css)
    writer = fitz.DocumentWriter(str(output))
    try:
        story.write(writer, lambda _index, _filled: (paper, area, None))
    finally:
        writer.close()
    with fitz.open(output) as document:
        if document.page_count < 1: raise ValueError("report_pdf_empty")
        for index, page in enumerate(document):
            page.insert_text((38, paper.height - 20), f"DeepEvo | {index+1}/{document.page_count}", fontsize=7, color=(0.3, 0.4, 0.38))
        document.saveIncr()


def render_document(source, root):
    root = Path(root)
    root.mkdir(parents=True, exist_ok=False)
    with _RENDER_LOCK:
        document = json.loads(canonical(source))
        document["document_sha256"] = __import__("hashlib").sha256(canonical(source).encode()).hexdigest()
        figures = render_figures(document["figures"], root / "figures")
        if any(figure["status"] != "ready" for figure in figures):
            document["report_status"] = "partial"
            document["missing"].extend(figure["id"] + ": " + figure.get("reason", "missing_data") for figure in figures if figure["status"] != "ready")
        document["rendered_figures"] = figures
        content = sections(document)
        (root / "report-document.json").write_text(canonical(document) + "\n", encoding="utf-8")
        (root / "report.md").write_text(render_markdown(document, content, figures), encoding="utf-8")
        html_text = render_html(document, content, figures, root)
        (root / "report.html").write_text(html_text, encoding="utf-8")
        if "docx" in document["formats"]: render_docx(document, content, figures, root)
        if "pdf" in document["formats"]: render_pdf(html_text, root / "report.pdf")
        if figures:
            from . import report_figures
            (root / "plot_figures.py").write_bytes(Path(report_figures.__file__).read_bytes())
            (root / "figure-sources.json").write_text(canonical(document["figures"]), encoding="utf-8")
        versions = {}
        for package in ("matplotlib", "python-docx", "PyMuPDF"):
            try: versions[package] = importlib.metadata.version(package)
            except importlib.metadata.PackageNotFoundError: versions[package] = None
        (root / "renderer-environment.json").write_text(canonical(versions), encoding="utf-8")
        files = [{"path": path.relative_to(root).as_posix(), "sha256": digest(path), "bytes": path.stat().st_size}
                 for path in sorted(root.rglob("*")) if path.is_file()]
        manifest = {"schema": "evomind.report_package.v2", "document_sha256": document["document_sha256"],
            "identity": document["identity"], "report_status": document["report_status"],
            "evidence_status": document["evidence_status"], "files": files,
            "hash_scope": "Payload only; manifest and ZIP excluded from recursive self-hashing."}
        (root / "report-manifest.json").write_text(canonical(manifest), encoding="utf-8")
        with zipfile.ZipFile(root / "report-bundle.zip", "x", compression=zipfile.ZIP_DEFLATED) as archive:
            for path in sorted(root.rglob("*")):
                if path.is_file() and path.name != "report-bundle.zip": archive.write(path, path.relative_to(root).as_posix())
        return manifest
