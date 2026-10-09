"""Real-model, bounded service soak in an isolated runtime; never uses GPU tools."""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import sys
import threading
import time
from urllib.parse import urlsplit


def write_json(path, value):
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=True, indent=2), encoding="utf-8")
    temporary.replace(path)


def verify_case_tool_evidence(calls, output_sha256, task_root):
    """Calling a tool is not proof that its requested operation succeeded."""
    root = Path(task_root).resolve(strict=True)
    expected_readback = root / 'outputs' / 'summary.json'

    def is_bound_readback(call):
        if call['tool_name'] != 'file_read':
            return False
        value = call.get('arguments', {}).get('path')
        if not isinstance(value, str) or not value.strip():
            return False
        try:
            candidate = Path(value)
            if not candidate.is_absolute():
                candidate = root / candidate
            # Store persists policy-normalized absolute paths; legacy relative
            # receipts are accepted only when they resolve to this exact Run.
            return candidate.resolve() == expected_readback
        except (OSError, RuntimeError, ValueError):
            return False

    completed = [call for call in calls if call.get('status') == 'completed' and (call.get('result') or {}).get('ok') is True]
    names = {call['tool_name'] for call in completed}
    hash_ok = bool(output_sha256) and any(
        call['tool_name'] == 'directory_hash'
        and call.get('arguments', {}).get('relative_path') == 'outputs/summary.json'
        and (call['result'].get('content') or {}).get('receipt', {}).get('sha256') == output_sha256
        and (call['result'].get('content') or {}).get('receipt', {}).get('ok') is True
        for call in completed)
    readback_ok = any(is_bound_readback(call) for call in completed)
    return {'successful_tool_count': len(completed), 'successful_tool_names': sorted(names),
            'hash_check_passed': hash_ok, 'readback_passed': readback_ok,
            'required_tool_steps_passed': hash_ok and readback_ok and {'file_read', 'file_write', 'directory_hash', 'artifact_publish'}.issubset(names)}


def verify_report_job_evidence(runtime, run_id, calls, output_sha256):
    """Bind report acceptance to a native request, durable job and readable bytes."""
    import io
    import re
    import zipfile
    from pathlib import PurePosixPath
    from evomind_runtime.report_document import canonical, digest

    completed = [call for call in calls if call.get('session_id') == run_id
                 and call.get('status') == 'completed' and (call.get('result') or {}).get('ok') is True]
    required_formats = {'markdown', 'html', 'docx', 'pdf'}
    origins = [call for call in completed if call.get('tool_name') == 'report_generate'
               and not call.get('arguments', {}).get('report_id')
               and required_formats.issubset(call.get('arguments', {}).get('formats') or [])]
    if not origins:
        return {'passed': False, 'error_code': 'report_generation_receipt_missing'}
    root = Path(runtime.get_session(run_id)['workspace_root']).resolve(strict=True)
    rejected = []
    checked_bytes = 0

    def require(condition, code):
        if not condition:
            raise ValueError(code)

    def checked_artifact(item):
        nonlocal checked_bytes
        stored = runtime.store.get_deliverable(item.get('id', ''))
        require(stored is not None and stored.get('run_id') == run_id
                and stored.get('session_id') == run_id, 'report_artifact_wrong_run')
        require(stored.get('source_tool_call') == 'report_generate', 'report_artifact_wrong_source')
        require(all(item.get(key) == stored.get(key) for key in ('name', 'sha256', 'bytes')),
                'report_artifact_receipt_drift')
        path = Path(stored['path'])
        path.resolve(strict=True).relative_to(root)
        require(not path.is_symlink() and path.is_file() and 0 < path.stat().st_size <= 32 * 1024 * 1024,
                'report_artifact_invalid')
        checked_bytes += path.stat().st_size
        require(checked_bytes <= 64 * 1024 * 1024, 'report_validation_size_limit')
        with path.open('rb') as handle:
            data = handle.read(32 * 1024 * 1024 + 1)
        require(len(data) == stored['bytes'] and hashlib.sha256(data).hexdigest() == stored['sha256'],
                'report_artifact_hash_mismatch')
        return stored, data

    for origin in reversed(origins):
        identifier = ''
        checked_bytes = 0
        try:
            receipt = origin['result']['content']['report_job']
            identifier = receipt.get('id', '')
            require(re.fullmatch(r'report_[a-f0-9]{32}', identifier) is not None
                    and receipt.get('run_id') == run_id, 'report_request_identity_invalid')
            job = runtime.reports.get(run_id, identifier)
            require(job['run_id'] == run_id and job['id'] == identifier, 'report_job_wrong_run')
            require(job['status'] in {'ready', 'partial'} and job.get('stage') == 'finished'
                    and not job.get('error_code'), 'report_job_not_terminal')
            document = job['document']
            document_sha = hashlib.sha256(canonical(document).encode()).hexdigest()
            require(job['document_sha256'] == document_sha == receipt.get('document_sha256')
                    and identifier == 'report_' + document_sha[:32], 'report_document_identity_mismatch')
            require(document.get('schema') == 'evomind.report_document.v1'
                    and document.get('identity', {}).get('run_id') == run_id
                    and required_formats.issubset(document.get('formats') or []), 'report_document_source_invalid')
            require(document.get('renderer_identity') == runtime.reports.renderer_identity(), 'report_renderer_identity_mismatch')
            sources = [source for source in document.get('sources', [])
                       if source.get('name') == 'summary.json' and source.get('sha256') == output_sha256
                       and source.get('verification') == 'hash_verified']
            require(bool(output_sha256) and bool(sources), 'report_current_summary_source_missing')
            for source in sources:
                artifact = runtime.store.get_deliverable(source['id'])
                require(artifact is not None and artifact.get('run_id') == run_id
                        and artifact.get('sha256') == output_sha256, 'report_summary_source_wrong_run')
                source_path = Path(artifact['path'])
                source_path.resolve(strict=True).relative_to(root)
                require(not source_path.is_symlink() and source_path.stat().st_size == source['bytes']
                        and digest(source_path) == output_sha256, 'report_summary_source_hash_mismatch')
            confirmations = [call for call in completed if call.get('tool_name') == 'report_status'
                             and call.get('arguments', {}).get('report_id') == identifier]
            require(any((call['result'].get('content') or {}).get('report_job', {}).get('id') == identifier
                        and call['result']['content']['report_job'].get('run_id') == run_id
                        and call['result']['content']['report_job'].get('document_sha256') == document_sha
                        and call['result']['content']['report_job'].get('status') == job['status']
                        for call in confirmations), 'report_terminal_receipt_missing')
            require(1 <= len(job.get('artifacts', [])) <= 512, 'report_artifact_count_invalid')
            published = [checked_artifact(item) for item in job['artifacts']]
            manifests = [(stored, data) for stored, data in published if stored['name'] == 'report-manifest.json']
            require(len(manifests) == 1, 'report_manifest_artifact_missing')
            manifest = json.loads(manifests[0][1])
            require(manifest == job.get('manifest') and manifest.get('schema') == 'evomind.report_package.v2'
                    and manifest.get('document_sha256') == document_sha
                    and manifest.get('identity') == document['identity'], 'report_manifest_source_mismatch')
            require(manifest.get('report_status') == job.get('report_status') == job['status']
                    and manifest.get('evidence_status') == document.get('evidence_status'), 'report_manifest_status_mismatch')
            files = manifest.get('files')
            require(isinstance(files, list) and 1 <= len(files) <= 512, 'report_manifest_files_invalid')
            payloads = {}
            for entry in files:
                relative = entry.get('path')
                require(isinstance(relative, str) and '\\' not in relative and ':' not in relative,
                        'report_manifest_path_invalid')
                name = PurePosixPath(relative)
                require(not name.is_absolute() and '..' not in name.parts and str(name) == relative
                        and relative not in payloads, 'report_manifest_path_invalid')
                matches = [(stored, data) for stored, data in published if stored['name'] == name.name
                           and stored['sha256'] == entry.get('sha256') and stored['bytes'] == entry.get('bytes')]
                require(len(matches) == 1, 'report_payload_not_published')
                payloads[relative] = matches[0][1]
            require({'report-document.json', 'report.md', 'report.html', 'report.docx', 'report.pdf'}.issubset(payloads),
                    'report_required_payload_missing')
            rendered = json.loads(payloads['report-document.json'])
            require(rendered.get('document_sha256') == document_sha
                    and rendered.get('report_status') == manifest['report_status']
                    and all(rendered.get(key) == document.get(key) for key in
                            ('schema', 'identity', 'model', 'formats', 'sources', 'metrics', 'checks', 'title', 'language',
                             'kind', 'execution_status', 'evidence_status', 'limitations', 'author_summary', 'renderer_identity')),
                    'report_rendered_document_mismatch')
            with zipfile.ZipFile(io.BytesIO(payloads['report.docx'])) as archive:
                members = archive.infolist()
                require(len(members) <= 1000 and sum(item.file_size for item in members) <= 32 * 1024 * 1024
                        and {'[Content_Types].xml', 'word/document.xml', '_rels/.rels'}.issubset(archive.namelist()),
                        'report_docx_structure_invalid')
            from docx import Document
            import fitz
            word = Document(io.BytesIO(payloads['report.docx']))
            word_text = '\n'.join([paragraph.text for paragraph in word.paragraphs]
                                  + [cell.text for table in word.tables for row in table.rows for cell in row.cells])
            require(run_id in word_text and document_sha in word_text and output_sha256 in word_text,
                    'report_docx_content_identity_mismatch')
            with fitz.open(stream=payloads['report.pdf'], filetype='pdf') as pdf:
                require(pdf.is_pdf and not pdf.needs_pass and pdf.page_count > 0, 'report_pdf_structure_invalid')
                pdf_pages = pdf.page_count
                # MuPDF preserves typographic ligatures by default: a genuine
                # ASCII hash containing "ff" is extracted as U+FB00. Expand
                # those glyphs at extraction time, not via broad Unicode
                # normalization that could accept a different identifier.
                text_flags = fitz.TEXTFLAGS_TEXT & ~fitz.TEXT_PRESERVE_LIGATURES
                pdf_text = ''.join(page.get_text(flags=text_flags) for page in pdf)
                require(run_id in pdf_text and output_sha256 in re.sub(r'\s+', '', pdf_text),
                        'report_pdf_content_identity_mismatch')
            return {'passed': True, 'report_id': identifier, 'document_sha256': document_sha,
                    'job_status': job['status'], 'manifest_sha256': manifests[0][0]['sha256'],
                    'checked_payload_files': len(payloads), 'docx_tables': len(word.tables), 'pdf_pages': pdf_pages}
        except ImportError:
            rejected.append({'report_id': identifier, 'error_code': 'report_validator_dependency_missing'})
        except Exception as error:
            code = str(error) if isinstance(error, ValueError) and re.fullmatch(r'report_[a-z_]+', str(error)) else 'report_evidence_invalid'
            rejected.append({'report_id': identifier, 'error_code': code})
    return {'passed': False, 'error_code': rejected[-1]['error_code'], 'rejected_jobs': rejected}


def bind_case_instructions(runtime, run, scoped_prompt):
    # Original Run prompts are immutable: update_assistant_run deliberately does
    # not accept a `prompt` field. Use the supported continuation contract and
    # verify the actual model-facing execution prompt before any model request.
    plan = {**run.get('plan', {}), 'continuation_instructions': [scoped_prompt]}
    updated = runtime.store.update_assistant_run(run['id'], plan=plan)
    rendered = runtime.assistant._execution_prompt(updated, resume=False)
    if scoped_prompt not in rendered:
        raise RuntimeError('acceptance_instructions_not_model_visible')
    return updated


def wait_for_case_slot(started, completed_cases, interval_seconds, duration_seconds, report, state_path):
    """Spread real tasks across the full observation window at bounded load."""
    if not interval_seconds:
        return
    target = min(started + completed_cases * interval_seconds, started + duration_seconds)
    while time.monotonic() < target:
        report.update(current_stage='waiting_next_case', elapsed_seconds=time.monotonic() - started)
        write_json(state_path, report)
        time.sleep(min(30, max(0, target - time.monotonic())))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--model", choices=["gpt-6-astra", "gpt-5.6-sol", 'gpt-5.6-terra', 'gpt-5.6-luna', 'gpt-5.5'], required=True)
    parser.add_argument("--seconds", type=int, default=5400)
    parser.add_argument("--root", required=True)
    parser.add_argument("--gateway-config", required=True)
    parser.add_argument("--candidate-manifest", required=True)
    parser.add_argument('--protocol', choices=['chat_completions', 'responses'], default='chat_completions')
    parser.add_argument('--tier', choices=['priority', 'omit'], default='priority')
    parser.add_argument('--preflight-cases', type=int, choices=[1, 2, 3])
    parser.add_argument('--route', choices=['gateway', 'configured_upstream'], default='gateway')
    parser.add_argument('--case-interval-seconds', type=int, choices=[0, 180], default=0)
    parser.add_argument('--report-every', type=int, choices=[1, 5], default=5)
    args = parser.parse_args()
    if args.seconds != 5400:
        raise SystemExit("formal_endurance_requires_5400_seconds")
    extras = args.model not in {'gpt-6-astra', 'gpt-5.6-sol', 'gpt-5.5'}
    if extras and (not args.preflight_cases or args.protocol != 'responses'):
        raise SystemExit('extra_candidates_bounded_responses_preflight_only')
    root = Path(args.root).resolve()
    root.mkdir(parents=True, exist_ok=False)
    manifest_path = Path(args.candidate_manifest).resolve()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    candidate_root = manifest_path.parent
    for row in manifest["files"]:
        path = (candidate_root / row["path"]).resolve()
        path.relative_to(candidate_root)
        if hashlib.sha256(path.read_bytes()).hexdigest() != row["sha256"]:
            raise SystemExit("candidate_integrity_failed")
    import yaml
    data = yaml.safe_load(Path(args.gateway_config).read_text(encoding="utf-8-sig"))
    keys = data.get("api-keys") or []
    if len(keys) != 1: raise SystemExit("gateway_client_binding_unavailable")
    from research_os.agent import messaging
    from research_os.llm_client import ProviderConfig
    original_client = messaging.AgentMessageClient
    os.environ['EVOMIND_MODEL_WIRE_PROTOCOL'] = args.protocol
    os.environ['EVOMIND_MODEL_TIMEOUT_SECONDS'] = '90'
    os.environ['EVOMIND_MODEL_ROUTE_CONFIG_PATH'] = str(Path(args.gateway_config).resolve())
    if extras:
        os.environ['EVOMIND_MODEL_ACCEPTANCE_EXTRAS'] = '1'
    endpoint, secret = 'http://127.0.0.1:65068/v1', keys[0]
    if args.route == 'configured_upstream':
        providers = [item for item in data.get('openai-compatibility', []) if item.get('name') == 'pezayo']
        if len(providers) != 1 or len(providers[0].get('api-key-entries', [])) != 1:
            raise SystemExit('configured_provider_ambiguous')
        entry = providers[0]['api-key-entries'][0]
        endpoint, secret = providers[0]['base-url'], entry['api-key']
        if urlsplit(endpoint).scheme != 'https' or urlsplit(endpoint).hostname != 'api.pezayo.com':
            raise SystemExit('unapproved_provider_origin')
        if str(entry.get('proxy-url', '')).lower() not in {'', 'direct', 'direct://'}:
            raise SystemExit('configured_proxy_requires_explicit_transport')
        # Honor the configured direct route in this isolated child process only.
        for variable in ('HTTP_PROXY', 'HTTPS_PROXY', 'ALL_PROXY', 'http_proxy', 'https_proxy', 'all_proxy'):
            os.environ.pop(variable, None)
        os.environ['NO_PROXY'] = os.environ['no_proxy'] = '*'
    config = ProviderConfig("openai", endpoint, args.model, secret, reasoning_effort="low", service_tier='priority' if args.tier == 'priority' else None)
    messaging.AgentMessageClient = lambda **kwargs: original_client(
        max_retries=0 if args.preflight_cases else kwargs.get("max_retries", 2), timeout=90, transports=[messaging.OpenAITransport(config)])
    from evomind_runtime.runtime import AgentRuntime
    runtime = AgentRuntime(root / "isolated-workspace")
    report = {"schema": "evomind.model_service_endurance.v1", "model": args.model, "status": "running",
        "harness_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "started_at": datetime.now(timezone.utc).isoformat(), "required_seconds": 5400,
        "candidate_manifest_sha256": hashlib.sha256(manifest_path.read_bytes()).hexdigest(),
        "same_entrypoint": "AgentRuntime.assistant._execute -> AgentRuntime.message -> governed_client",
        "provider_target": urlsplit(endpoint).netloc, "route": args.route, "reasoning_effort": "low", "service_tier": args.tier,
        "wire_protocol": args.protocol, "preflight_only": bool(args.preflight_cases),
        "case_interval_seconds": args.case_interval_seconds, "load_scope": "paced_service_endurance" if args.case_interval_seconds else "continuous_service_endurance",
        "report_every": args.report_every,
        "cases": [], "native_tools_executed": 0, "real_transport_failures": 0, "injected_faults": 0,
        "native_tool_records": 0, "failed_tool_records": 0,
        "gpu_actions": 0, "production_default_changed": False, "qualified": False}
    report.update(service_soak_passed=False, final_model_qualification='pending_hard_gate_evidence',
                  hard_gates_not_proved_by_this_harness=['fault_injection', 'cross_owner_authorization',
                    'browser_disconnect', 'service_restart_recovery', 'approval_replay', 'http_ack_p95'])
    started = time.monotonic()
    state_path = root / "endurance.json"
    write_json(state_path, report)
    print(json.dumps({"model": args.model, "status": "started", "root": str(root)}), flush=True)
    try:
        case_index = 0
        while time.monotonic() - started < args.seconds:
            if case_index and not args.preflight_cases:
                wait_for_case_slot(started, case_index, args.case_interval_seconds, args.seconds, report, state_path)
                if time.monotonic() - started >= args.seconds:
                    continue
            case_index += 1
            report.update(current_stage='executing_case', current_case=case_index,
                          elapsed_seconds=time.monotonic() - started)
            write_json(state_path, report)
            values = [case_index, case_index + 1, case_index + 2, -case_index]
            expected = {"case_id": case_index, "count": 4, "sum": sum(values), "min": min(values), "max": max(values)}
            needs_report = case_index % args.report_every == 0
            prompt = (
                f"Isolated system acceptance case {case_index}; synthetic data, not research or competition evaluation. "
                "Read inputs/data.json with file_read. Write outputs/summary.json containing exactly the JSON keys "
                "case_id, count, sum, min, max computed from its values. Read outputs/summary.json back with file_read. "
                "Use directory_hash on the file relative_path outputs/summary.json (not the outputs directory), "
                "and artifact_publish the summary. Do not use shell, network, GPU, training, downloads, credentials, or unrelated paths. "
                + ("Report dependency: first finish artifact_publish(outputs/summary.json) and receive its successful artifact receipt. "
                   "Only AFTER that tool result, call report_generate with artifact_ids=[the exact returned summary artifact id], "
                   "title 'System acceptance fixture', report_kind 'analysis', language 'zh-CN', "
                   "formats ['markdown','html','docx','pdf']. Do not put report generation in the same batch as publication. "
                   "Query report_status until ready or partial, verify the source IDs include that summary artifact, "
                   "and confirm report.docx/report.pdf were published. A report created before source publication does not meet this task. " if needs_report else "")
                + "Finish with a concise factual summary. Do not label the fixture as a scientific success."
            )
            run = runtime.assistant.create_run(prompt=prompt, conversation_id=f"endurance_{args.model.replace('-','_')}", start=False)
            session = runtime.get_session(run["id"])
            directories = session["metadata"].get("super_agent_directory_ids") or []
            if not directories:
                raise RuntimeError("acceptance_directory_binding_missing")
            scoped_prompt = prompt + " For directory_hash use exactly " + json.dumps({
                'directory_id': directories[0], 'relative_path': 'outputs/summary.json'}) + ". Do not invent a directory_id. A rejected/failed hash call is not a successful verification."
            bind_case_instructions(runtime, run, scoped_prompt)
            allowed = ["file_read", "file_write", "directory_hash", "artifact_publish", "artifact_list", "report_generate", "report_status"]
            runtime.store.update_session(run["id"], objective=scoped_prompt, metadata_json={**session["metadata"], "run_allowed_tool_names": allowed,
                                                                 "acceptance_scope": "synthetic_system_soak"})
            task_root = Path(session["workspace_root"])
            (task_root / "inputs").mkdir(exist_ok=True)
            (task_root / "inputs/data.json").write_text(json.dumps({"case_id": case_index, "values": values}), encoding="utf-8")
            begin = time.monotonic()
            runtime.assistant.start(run["id"])
            worker = runtime.assistant._threads.get(run["id"])
            if worker:
                worker.join(600)  # finish a case begun before the 90-minute window ends
            timed_out = bool(worker and worker.is_alive())
            if timed_out:
                runtime.assistant.action(run["id"], "pause")
                worker.join(150)  # allow the one already-in-flight request to settle
            snapshot = runtime.assistant.snapshot(run["id"])
            calls = runtime.store.list_tool_calls(run["id"], limit=2000)
            events = runtime.store.list_events(run["id"])
            transport = [event["payload"] for event in events if event["event_type"] == "model.transport_attempt"]
            failures = [event for event in transport if event.get("status") == "failed"]
            tools = [call for call in calls if call["status"] in {"completed", "failed"}]
            output = task_root / "outputs/summary.json"
            actual = json.loads(output.read_text(encoding="utf-8")) if output.is_file() else None
            names = {item["name"] for item in runtime.store.list_deliverables(run["id"])}
            output_digest = hashlib.sha256(output.read_bytes()).hexdigest() if output.is_file() else None
            evidence = verify_case_tool_evidence(calls, output_digest, task_root)
            success = not timed_out and snapshot["status"] == "completed" and actual == expected and "summary.json" in names
            success = success and evidence['required_tool_steps_passed']
            if needs_report:
                evidence['report_job_evidence'] = verify_report_job_evidence(runtime, run['id'], calls, output_digest)
                success = success and evidence['report_job_evidence']['passed']
            row = {"index": case_index, "run_id": run["id"], "status": snapshot["status"], "passed": success,
                "elapsed_seconds": time.monotonic() - begin, "tool_count": len(tools), "transport_attempts": len(transport),
                "transport_failures": len(failures), "failure_codes": [event.get("error_code") for event in failures],
                "needs_report": needs_report, "artifact_count": len(names),
                "case_timeout": timed_out,
                "output_sha256": output_digest, "tool_evidence": evidence}
            report["cases"].append(row)
            report["native_tool_records"] += len(tools)
            report["failed_tool_records"] += sum(call['status'] == 'failed' for call in tools)
            report["native_tools_executed"] += evidence['successful_tool_count']
            report["real_transport_failures"] += len(failures)
            report["elapsed_seconds"] = time.monotonic() - started
            write_json(state_path, report)
            print(json.dumps({"model": args.model, **row}, ensure_ascii=True), flush=True)
            if not success:
                report.update(status="failed", failure_code="functional_case_failed")
                break
            if args.preflight_cases and case_index >= args.preflight_cases:
                report.update(status='preflight_passed')
                break
        else:
            report["status"] = "completed"
        report["elapsed_seconds"] = time.monotonic() - started
        report["service_soak_passed"] = not args.preflight_cases and report["status"] == "completed" and report["elapsed_seconds"] >= 5400 and len(report["cases"]) >= 20 and report["native_tools_executed"] >= 50
    except Exception as error:
        report.update(status="failed", failure_code=type(error).__name__, qualified=False, elapsed_seconds=time.monotonic()-started)
    except KeyboardInterrupt:
        report.update(status='cancelled', failure_code='operator_interrupted', qualified=False,
                      elapsed_seconds=time.monotonic()-started)
    finally:
        report["runtime_closed_cleanly"] = runtime.close(timeout=30)
        if not report["runtime_closed_cleanly"]: report["service_soak_passed"] = False
        report["completed_at"] = datetime.now(timezone.utc).isoformat()
        write_json(state_path, report)
        print(json.dumps({"model": args.model, "status": report["status"], "qualified": report["qualified"],
                          "elapsed_seconds": report["elapsed_seconds"], "cases": len(report["cases"]),
                          "native_tools_executed": report["native_tools_executed"]}), flush=True)
    return 0 if report["service_soak_passed"] or report['status'] == 'preflight_passed' else 2


if __name__ == "__main__":
    raise SystemExit(main())
