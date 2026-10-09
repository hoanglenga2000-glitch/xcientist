"""Install and prove a delivered report wheelhouse without activating it.

Only the named C:/EMQA/*-report-delivered-vN target is writable. Installation
and rendering are separate, single-use phases; existing receipts require
inspection instead of an automatic retry. No model or HPC action is needed.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
from email.parser import BytesParser
import hashlib
import importlib
import importlib.metadata
import importlib.machinery
import json
import os
from pathlib import Path, PurePosixPath
import platform
import re
import stat
import struct
import subprocess
import sys
import sysconfig
import time
import uuid
import zipfile


BASE_VENV = Path('C:/EvoMind/releases/664a636ddd419a66f73cc10820c0a16429784866/.venv')
BASE_TREE = '3800faca9fc507c17152904af09629ccc98a74769723e4f7d1ae02ed6f1086be'
BASE_MANIFEST = 'e3d80e0abd54dd73b4777dbe7b2b4f62bf488f9e4355119d297eb7931dab9507'
MODULES = {'contourpy':'contourpy', 'cycler':'cycler', 'fonttools':'fontTools', 'kiwisolver':'kiwisolver',
           'matplotlib':'matplotlib', 'numpy':'numpy', 'packaging':'packaging', 'pillow':'PIL',
           'pyparsing':'pyparsing', 'python-dateutil':'dateutil', 'six':'six',
           'python-docx':'docx', 'pymupdf':'fitz', 'lxml':'lxml', 'typing-extensions':'typing_extensions'}


def require(value, code):
    if not value:
        raise ValueError(code)


def normalized(name):
    return re.sub(r'[-_.]+', '-', name).lower()


def canonical(value):
    return json.dumps(value, sort_keys=True, ensure_ascii=True, separators=(',', ':'), allow_nan=False)


def digest(path):
    result = hashlib.sha256()
    with Path(path).open('rb') as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b''):
            result.update(block)
    return result.hexdigest()


def safe_relative(value):
    require(isinstance(value, str) and bool(value) and '\\' not in value and ':' not in value,
            'relative_path_rejected')
    path = PurePosixPath(value)
    require(not path.is_absolute() and '..' not in path.parts and '.' not in path.parts
            and str(path) == value, 'relative_path_rejected')
    return path


def checked_file(root, row):
    relative = safe_relative(row['path'])
    path = root.joinpath(*relative.parts)
    path.resolve(strict=True).relative_to(root.resolve(strict=True))
    require(not path.is_symlink() and path.is_file() and path.stat().st_size == row['bytes']
            and digest(path) == row['sha256'], 'declared_file_integrity_failed')
    return path


def checked_json(path, expected):
    require(re.fullmatch(r'[a-f0-9]{64}', expected) is not None and digest(path) == expected,
            'input_receipt_hash_mismatch')
    return json.loads(path.read_text(encoding='utf-8'))


def write_json(path, value):
    temporary = path.with_name(path.name + '.' + uuid.uuid4().hex + '.tmp')
    with temporary.open('x', encoding='utf-8') as handle:
        handle.write(canonical(value) + '\n')
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temporary, path)


def tree_manifest(root):
    rows = []
    for path in sorted(root.rglob('*')):
        require(not path.is_symlink(), 'manifest_symlink_rejected')
        if path.is_file():
            rows.append({'path':path.relative_to(root).as_posix(), 'bytes':path.stat().st_size, 'sha256':digest(path)})
    return {'files':rows, 'file_count':len(rows), 'tree_sha256':hashlib.sha256(canonical(rows).encode()).hexdigest()}


def verify_rows(root, manifest):
    paths = [row['path'] for row in manifest['files']]
    require(len(paths) == len(set(paths)), 'duplicate_manifest_path')
    for row in manifest['files']:
        checked_file(root, row)
    return len(paths)


def validate_locations(stage, target):
    stage = Path(stage).absolute()
    target = Path(target).absolute()
    require(stage.parent == Path('C:/ProgramData/EvoMind/staging')
            and re.fullmatch(r'invitation-[a-f0-9]{12}-sys\d+', stage.name), 'stage_scope_rejected')
    require(target.parent == Path('C:/EMQA')
            and re.fullmatch(r'[A-Za-z0-9-]+-report-delivered-v\d+', target.name), 'target_scope_rejected')
    require(stage.resolve(strict=True) == stage and not stage.is_symlink()
            and target.resolve(strict=False) == target and not target.is_symlink(), 'location_alias_rejected')
    return stage, target


def validate_wheel(path, expected_name, expected_version, supported_tags):
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        require(len(members) <= 10000 and sum(item.file_size for item in members) <= 512*1024*1024,
                'wheel_expansion_limit')
        for item in members:
            safe_relative(item.filename.rstrip('/'))
            require(not stat.S_ISLNK(item.external_attr >> 16) and not item.filename.endswith('.pth'),
                    'wheel_executable_path_rejected')
        metadata = [item for item in members if item.filename.endswith('.dist-info/METADATA')]
        wheel = [item for item in members if item.filename.endswith('.dist-info/WHEEL')]
        require(len(metadata) == len(wheel) == 1, 'wheel_metadata_missing')
        meta = BytesParser().parsebytes(archive.read(metadata[0]))
        tags = BytesParser().parsebytes(archive.read(wheel[0])).get_all('Tag', [])
        require(normalized(meta['Name']) == expected_name and meta['Version'] == expected_version,
                'wheel_distribution_mismatch')
        require(any(tag in supported_tags for tag in tags), 'wheel_abi_incompatible')


def delivered_runtime_manifest(stage, acceptance):
    # The outer acceptance manifest binds receipts, not extracted module files.
    # Follow its hash-bound source receipt before trusting a runtime import.
    rows = [row for row in acceptance['files'] if row['path']=='source-receipt.json']
    require(len(rows)==1, 'source_receipt_binding_missing')
    source = checked_json(stage/'source-receipt.json', rows[0]['sha256'])
    require(source.get('frozen') is True and source.get('source_tree_sha256')==acceptance.get('source_tree_sha256'),
            'source_receipt_identity_mismatch')
    runtime_rows = [row for row in source['files'] if row['path'].startswith('runtime/')]
    require(bool(runtime_rows) and len(runtime_rows)==source['runtime_file_count'], 'runtime_file_count_mismatch')
    verify_rows(stage, {'files':runtime_rows})
    return source


def preflight(args, stage):
    wheelhouse = stage / 'web/support/report-wheelhouse'
    descriptor = checked_json(wheelhouse/'dependency-delivery.json', args.descriptor_sha256)
    inheritance = checked_json(wheelhouse/'report-inheritance.json', args.inheritance_sha256)
    acceptance = checked_json(stage/'acceptance-manifest.json', args.acceptance_sha256)
    require(descriptor.get('schema') == 'evomind.syscomplete_report_dependencies.v1'
            and descriptor.get('frozen') is True and descriptor.get('python_tag') == 'cp312'
            and descriptor.get('platform') == 'win_amd64', 'descriptor_contract_invalid')
    require(descriptor['base_environment_binding']['receipt_sha256'] == args.inheritance_sha256
            and descriptor['base_environment_binding']['base_build_id'] == inheritance['base_build_id']
            and inheritance.get('base_source_tree_sha256') == BASE_TREE
            and inheritance.get('base_source_manifest_sha256') == BASE_MANIFEST
            and inheritance.get('production_changed') is False and inheritance.get('activation_claim') is False,
            'inheritance_binding_invalid')
    require(Path(inheritance['venv_root']).resolve() == BASE_VENV.resolve() == Path(sys.prefix).resolve(),
            'base_python_binding_invalid')
    abi = {'python_version':platform.python_version(), 'platform':sysconfig.get_platform(),
           'python_implementation':platform.python_implementation(), 'cache_tag':sys.implementation.cache_tag,
           'pointer_bits':struct.calcsize('P')*8, 'soabi':sysconfig.get_config_var('SOABI'),
           'extension_suffixes':importlib.machinery.EXTENSION_SUFFIXES}
    require(abi['python_version'] == inheritance['python_version'] and sys.version_info[:2] == (3,12)
            and abi['platform'].replace('-', '_') == 'win_amd64' and abi['cache_tag'] == 'cpython-312'
            and abi['python_implementation'] == 'CPython' and abi['pointer_bits'] == 64
            and abi['soabi'] == inheritance.get('python_abi'), 'base_python_abi_mismatch')
    expected = {normalized(row['name']):row for row in descriptor['distributions']}
    require(set(expected) == set(MODULES) and len(expected) == len(descriptor['distributions']),
            'distribution_allowlist_mismatch')
    inherited = {normalized(row['name']):row['version'] for row in inheritance['packages']}
    require(inherited == {name:row['version'] for name,row in expected.items() if row['source']=='pinned_base_environment'},
            'inherited_versions_mismatch')
    count = 0
    for package in inheritance['packages']:
        require(package.get('import_passed') is True, 'inherited_import_receipt_invalid')
        for row in package['files']:
            require(row['path'].startswith('Lib/site-packages/')
                    or normalized(package['name']) == 'pymupdf' and row['path'] == 'Scripts/pymupdf.exe',
                    'inherited_file_scope_rejected')
            checked_file(BASE_VENV, row)
            count += 1
        distribution = importlib.metadata.distribution(package['name'])
        require(distribution.version == package['version'], 'inherited_live_version_mismatch')
        imported = importlib.import_module(package['module'])
        actual = Path(imported.__file__).resolve()
        require(actual == Path(package['module_path']).resolve()
                and actual in {BASE_VENV.joinpath(*safe_relative(row['path']).parts).resolve() for row in package['files']},
                'inherited_import_origin_mismatch')
    require(count == 573, 'inherited_record_count_mismatch')
    verify_rows(stage, acceptance)
    require(digest(stage/'runtime.zip')==acceptance['runtime_sha256']
            and digest(stage/'web.zip')==acceptance['web_sha256'], 'delivered_archive_hash_mismatch')
    source = delivered_runtime_manifest(stage, acceptance)
    from pip._vendor.packaging.tags import sys_tags
    supported_tags = {str(tag) for tag in sys_tags()}
    wheels = descriptor['wheels']
    require(len(wheels) == 11 and {p.name for p in wheelhouse.glob('*.whl')} == {row['path'] for row in wheels},
            'delivered_wheel_set_mismatch')
    seen = set()
    for row in wheels:
        path = checked_file(wheelhouse, row)
        name, version = path.name.split('-')[:2]
        name = normalized(name)
        require(name not in seen and expected[name] == {'name':expected[name]['name'], 'version':version, 'source':'wheelhouse'},
                'wheel_descriptor_mismatch')
        validate_wheel(path, name, version, supported_tags)
        seen.add(name)
    return descriptor, inheritance, source, {'abi':abi, 'inherited_file_count':count,
        'acceptance_files_verified':len(acceptance['files']), 'wheel_count':len(wheels),
        'delivered_runtime_files_verified':source['runtime_file_count'], 'source_receipt_sha256':digest(stage/'source-receipt.json'),
        'base_soabi_note':'Windows SOABI may be null; CPython cache tag, architecture, platform and actual wheel tags are also verified.'}


def configure_paths(stage, target):
    prefixes = [stage/'runtime', target/'site-packages', BASE_VENV/'Lib/site-packages',
                stage/'web/support/python-runtime', stage/'web/support/python-vendor', BASE_VENV.parent/'src']
    for path in prefixes:
        require(path.is_dir(), 'pythonpath_binding_missing')
    selected = [str(path.resolve()) for path in prefixes]
    sys.path[:] = selected + [path for path in sys.path if path not in selected]
    importlib.invalidate_caches()
    return selected


def verify_imports(descriptor, target):
    results = []
    for row in descriptor['distributions']:
        name = normalized(row['name'])
        imported = importlib.import_module(MODULES[name])
        distribution = importlib.metadata.distribution(row['name'])
        path = Path(imported.__file__).resolve()
        expected_root = target/'site-packages' if row['source']=='wheelhouse' else BASE_VENV/'Lib/site-packages'
        path.relative_to(expected_root.resolve())
        require(distribution.version == row['version'], 'installed_version_mismatch')
        results.append({'name':row['name'], 'version':distribution.version, 'source':row['source'],
                        'module':MODULES[name], 'module_path':str(path), 'module_sha256':digest(path)})
    return results


def verify_runtime_origins(stage, acceptance):
    expected = {row['path']:row for row in acceptance['files']}
    for name in ('runtime', 'managed_scheduler', 'report_document', 'report_jobs', 'report_render', 'report_figures'):
        importlib.import_module('evomind_runtime.' + name)
    records = []
    for name, module in list(sys.modules.items()):
        if not name.startswith('evomind_runtime') or not getattr(module, '__file__', None):
            continue
        path = Path(module.__file__).resolve()
        path.relative_to((stage/'runtime/evomind_runtime').resolve())
        relative = path.relative_to(stage).as_posix()
        require(relative in expected and digest(path) == expected[relative]['sha256'], 'runtime_origin_not_delivered')
        records.append({'module':name, 'path':str(path), 'sha256':expected[relative]['sha256']})
    return sorted(records, key=lambda row:row['module'])


def render_control(stage, target, acceptance):
    # Every attempted network connection is denied before it reaches a socket.
    network_attempts = []
    def audit(event, arguments):
        if event in {'socket.connect', 'socket.getaddrinfo'}:
            network_attempts.append(event)
            raise RuntimeError('network_forbidden_report_delivery_acceptance')
    sys.addaudithook(audit)
    verify_runtime_origins(stage, acceptance)
    from evomind_runtime.runtime import AgentRuntime
    runtime = AgentRuntime(target/'fixture-workspace')
    try:
        run = runtime.assistant.create_run(prompt='Synthetic report dependency delivery check; CPU document fixture only', start=False)
        source = {'rmse':0.125, 'schema':'evomind.figure_source.v1', 'figures':[{
            'id':'recorded-fixture', 'kind':'line', 'title':'Recorded synthetic observations',
            'x':[0,1,2,3], 'y':[1,None,0.4,0.1], 'x_label':'Step', 'y_label':'Fixture loss',
            'caption':'Synthetic engineering fixture. Missing observations remain missing.'}]}
        for name, args in [('file_write', {'path':'outputs/summary.json', 'content':canonical(source)}),
                           ('file_read', {'path':'outputs/summary.json'})]:
            result = runtime.invoke_tool(run['id'], name, args)
            require(result['status']=='completed' and result['result']['ok'], 'fixture_source_tool_failed')
        result = runtime.invoke_tool(run['id'], 'artifact_publish', {'path':'outputs/summary.json'})
        require(result['status']=='completed' and result['result']['ok'], 'fixture_source_publish_failed')
        artifact = result['result']['content']['artifact']
        result = runtime.invoke_tool(run['id'], 'report_generate', {
            'title':'报告依赖交付验收（合成数据，非科研成果）', 'report_kind':'analysis', 'language':'zh-CN',
            'summary':'仅用于验证离线依赖、来源绑定与报告格式；未调用模型或执行训练。',
            'artifact_ids':[artifact['id']], 'formats':['markdown','html','docx','pdf']})
        require(result['status']=='completed' and result['result']['ok'], 'fixture_report_submit_failed')
        identifier = result['result']['content']['report_job']['id']
        deadline = time.monotonic()+150
        while True:
            job = runtime.reports.get(run['id'], identifier)
            if job['status'] in {'ready','partial','failed'}:
                break
            require(time.monotonic()<deadline, 'fixture_report_timeout')
            time.sleep(0.05)
        require(job['status']=='ready' and job['source_artifact_ids']==[artifact['id']], 'fixture_report_not_ready')
        result = runtime.invoke_tool(run['id'], 'report_status', {'report_id':identifier})
        require(result['result']['ok'], 'fixture_report_status_failed')
        published = {item['name']:runtime.store.get_deliverable(item['id']) for item in job['artifacts']}
        for item in published.values():
            require(item['run_id']==run['id'] and item['source_tool_call']=='report_generate'
                    and digest(Path(item['path']))==item['sha256'], 'report_published_hash_mismatch')
        manifest = json.loads(Path(published['report-manifest.json']['path']).read_text(encoding='utf-8'))
        require(manifest==job['manifest'] and manifest['document_sha256']==job['document_sha256'], 'report_manifest_mismatch')
        for row in manifest['files']:
            item = published[PurePosixPath(row['path']).name]
            require(item['sha256']==row['sha256'] and item['bytes']==row['bytes'], 'report_payload_mismatch')
        from docx import Document
        import fitz
        word = Document(published['report.docx']['path'])
        word_text = '\n'.join([p.text for p in word.paragraphs] + [c.text for t in word.tables for r in t.rows for c in r.cells])
        with fitz.open(published['report.pdf']['path']) as pdf:
            require(pdf.is_pdf and pdf.page_count>0, 'report_pdf_invalid')
            pdf_text = '\n'.join(page.get_text() for page in pdf)
            pages = pdf.page_count
        texts = [word_text, pdf_text, Path(published['report.md']['path']).read_text(encoding='utf-8'),
                 Path(published['report.html']['path']).read_text(encoding='utf-8')]
        require(bool(word.tables) and all(run['id'] in text and '0.125' in text for text in texts), 'report_content_mismatch')
        calls = runtime.store.list_tool_calls(run['id'])
        allowed = {'file_write','file_read','artifact_publish','report_generate','report_status'}
        require(all(call['tool_name'] in allowed for call in calls) and not network_attempts, 'fixture_external_action_detected')
        require(not any(event['event_type'].startswith('model.') for event in runtime.store.list_events(run['id'])), 'fixture_model_request_detected')
        origins = verify_runtime_origins(stage, acceptance)
        unmodified = [{'module':name, 'path':str(Path(module.__file__).resolve())}
                      for name,module in list(sys.modules.items()) if name in {'research_os', 'research_os.agent.messaging', 'xsci'}
                      and getattr(module, '__file__', None)]
        return {'run_id':run['id'], 'report_id':identifier, 'source_artifact_id':artifact['id'],
            'source_sha256':artifact['sha256'], 'document_sha256':job['document_sha256'],
            'manifest_sha256':published['report-manifest.json']['sha256'], 'pdf_pages':pages,
            'editable_word_tables':len(word.tables), 'model_requests':0, 'hpc_actions':0, 'network_attempts':0,
            'runtime_origins':origins, 'unmodified_dependency_origins':unmodified,
            'unmodified_dependency_scope':'Origin recorded only; not a new whole-code acceptance.',
            'artifacts':list(published.values())}
    finally:
        require(runtime.close(timeout=15), 'fixture_runtime_close_incomplete')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('phase', choices=['install', 'render'])
    parser.add_argument('--stage', required=True)
    parser.add_argument('--target', required=True)
    parser.add_argument('--descriptor-sha256', required=True)
    parser.add_argument('--inheritance-sha256', required=True)
    parser.add_argument('--acceptance-sha256', required=True)
    args = parser.parse_args()
    sys.dont_write_bytecode = True
    os.environ.update(PYTHONDONTWRITEBYTECODE='1', PYTHONUTF8='1', CUDA_VISIBLE_DEVICES='', NVIDIA_VISIBLE_DEVICES='',
                      PIP_CONFIG_FILE=os.devnull)
    stage, target = validate_locations(args.stage, args.target)
    if args.phase=='install':
        require(not target.exists(), 'existing_target_requires_pid_and_receipt_reconciliation')
        target.mkdir()
    else:
        require(target.is_dir() and not (target/'render-receipt.json').exists(), 'existing_render_requires_reconciliation')
    receipt_path = target/(args.phase+'-receipt.json')
    receipt = {'schema':'evomind.report_dependency_delivery_acceptance.v1', 'phase':args.phase, 'status':'running',
        'pid':os.getpid(), 'started_at':datetime.now(timezone.utc).isoformat(), 'python_executable':sys.executable,
        'script_sha256':digest(Path(__file__)), 'stage':str(stage), 'target':str(target/'site-packages'),
        'descriptor_sha256':args.descriptor_sha256, 'inheritance_sha256':args.inheritance_sha256,
        'acceptance_manifest_sha256':args.acceptance_sha256, 'activated':False, 'production_changed':False,
        'model_requests':0, 'hpc_actions':0}
    write_json(receipt_path, receipt)
    before = None
    try:
        for name in ('temp', 'cache', 'matplotlib-cache'):
            (target/name).mkdir(exist_ok=True)
        os.environ.update(TEMP=str(target/'temp'), TMP=str(target/'temp'), TMPDIR=str(target/'temp'),
                          MPLCONFIGDIR=str(target/'matplotlib-cache'), XDG_CACHE_HOME=str(target/'cache'))
        before = tree_manifest(BASE_VENV)
        write_json(target/(args.phase+'-base-before.json'), before)
        descriptor, inheritance, acceptance, checks = preflight(args, stage)
        receipt.update(preflight=checks, stage_progress='inputs_and_inheritance_verified')
        write_json(receipt_path, receipt)
        if args.phase=='install':
            wheelhouse = stage/'web/support/report-wheelhouse'
            requirements = '\n'.join((wheelhouse/row['path']).as_uri()+' --hash=sha256:'+row['sha256'] for row in descriptor['wheels'])+'\n'
            lock = target/'requirements.lock'
            with lock.open('x', encoding='utf-8') as handle: handle.write(requirements)
            command = [sys.executable, '-B', '-m', 'pip', '--isolated', '--disable-pip-version-check', 'install',
                '--no-index', '--no-deps', '--no-cache-dir', '--require-hashes', '--no-compile',
                '--only-binary=:all:', '--find-links', str(wheelhouse), '--target', str(target/'site-packages'), '-r', str(lock)]
            receipt.update(stage_progress='installing_exact_offline_wheels', requirements_sha256=digest(lock), pip_command=command)
            write_json(receipt_path, receipt)
            with (target/'pip.stdout.log').open('xb') as out, (target/'pip.stderr.log').open('xb') as err:
                result = subprocess.run(command, env=dict(os.environ), stdout=out, stderr=err, timeout=300)
            require(result.returncode==0, 'offline_pip_install_failed')
            installed = tree_manifest(target/'site-packages')
            write_json(target/'installed-files.json', installed)
            receipt.update(installed_manifest_sha256=digest(target/'installed-files.json'), installed_file_count=installed['file_count'])
        else:
            previous = json.loads((target/'install-receipt.json').read_text(encoding='utf-8'))
            require(previous['status']=='staged_not_activated' and previous['descriptor_sha256']==args.descriptor_sha256,
                    'successful_install_receipt_required')
            installed = checked_json(target/'installed-files.json', previous['installed_manifest_sha256'])
            require(tree_manifest(target/'site-packages')==installed, 'installed_dependency_drift')
            receipt['installed_manifest_sha256'] = previous['installed_manifest_sha256']
        receipt['pythonpath_bindings'] = configure_paths(stage, target)
        receipt['pythonpath_binding_mode'] = 'process_only_sys_path_prefix_no_production_configuration_change'
        receipt['imports'] = verify_imports(descriptor, target)
        receipt['stage_progress'] = 'imports_verified'
        write_json(receipt_path, receipt)
        if args.phase=='render':
            receipt['report'] = render_control(stage, target, acceptance)
        require(tree_manifest(target/'site-packages')==installed, 'dependency_changed_during_verification')
        receipt.update(status='staged_not_activated', stage_progress='complete')
    except Exception as error:
        detail = str(error)
        receipt.update(status='failed', error_type=type(error).__name__,
                       error_code=detail if re.fullmatch(r'[a-z_]{1,160}',detail) else 'dependency_delivery_phase_failed')
    finally:
        if before is not None:
            after = tree_manifest(BASE_VENV)
            write_json(target/(args.phase+'-base-after.json'), after)
            unchanged = before==after
            receipt.update(base_venv_unchanged=unchanged, base_file_count=after['file_count'],
                           base_tree_sha256_before=before['tree_sha256'], base_tree_sha256_after=after['tree_sha256'])
            if not unchanged:
                receipt.update(status='failed', error_code='base_environment_changed', production_changed=True)
        receipt['completed_at'] = datetime.now(timezone.utc).isoformat()
        write_json(receipt_path, receipt)
        print(canonical({'receipt':str(receipt_path), 'receipt_sha256':digest(receipt_path), 'status':receipt['status'],
                         'error_code':receipt.get('error_code'), 'base_venv_unchanged':receipt.get('base_venv_unchanged')}), flush=True)
    return 0 if receipt['status']=='staged_not_activated' else 2


if __name__=='__main__':
    raise SystemExit(main())
