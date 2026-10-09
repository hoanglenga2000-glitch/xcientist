"""Real rendering QA using a selected document runtime; no network or GPU."""
import json
from pathlib import Path
import sys

from evomind_runtime.report_render import render_document
from evomind_runtime.report_document import digest

root = Path(sys.argv[1])
source = {"schema": "evomind.report_document.v1", "language": "zh-CN", "kind": "training",
    "identity": {"run_id": "run_report_fixture", "project_id": "project_fixture", "task_id": ""},
    "title": "科研报告排版验收（合成数据，非科研成绩）", "execution_status": "completed",
    "evidence_status": "hash_verified", "report_status": "partial", "author_summary": "用于检查中文排版、数值一致性和可编辑表格。",
    "metrics": [{"name": "rmse", "value": 1.23456789, "unit": "target units", "direction": "lower", "source_artifact_id": "artifact_fixture"}],
    "checks": [], "sources": [{"id": "artifact_fixture", "name": "fixture.json", "bytes": 100, "sha256": "a"*64, "verification": "hash_verified"}],
    "missing": ["合成数据不代表真实模型训练验收。"], "limitations": ["Synthetic engineering fixture, not a scientific result."],
    "formats": ["markdown", "html", "docx", "pdf"],
    "figures": [{"kind": "line", "title": "Recorded loss fixture", "x": [0, 1, 2, 3], "y": [1, None, -0.2, 0.4],
                 "x_label": "Step", "y_label": "Loss", "source_artifact_id": "artifact_fixture", "source_sha256": "a"*64}]}
manifest = render_document(source, root)
from docx import Document
import fitz
word = Document(root / "report.docx")
word_text = "\n".join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
with fitz.open(root / "report.pdf") as pdf:
    pdf_text = "\n".join(page.get_text() for page in pdf)
    pages = pdf.page_count
print(json.dumps({"docx_editable_table": bool(word.tables), "pdf_pages": pages,
    "metric_in_all_formats": all("1.23456789" in value for value in (word_text, pdf_text, (root / "report.md").read_text(encoding="utf-8"), (root / "report.html").read_text(encoding="utf-8"))),
    "all_payload_hashes_verified": all(digest(root / row["path"]) == row["sha256"] for row in manifest["files"])}))
